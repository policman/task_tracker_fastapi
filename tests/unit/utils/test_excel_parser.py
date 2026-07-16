import openpyxl

from app.utils.excel_parser import parse_excel_generator


class TestExcelParser:

    def test_parse_excel_generator_happy_path(self, tmp_path):
        excel_file = tmp_path / "test_data.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active

        ws.append(["id", "name", "role"])
        ws.append([1, "admin", "superuser"])
        ws.append([2, "user", "manager"])

        wb.save(excel_file)

        result = list(parse_excel_generator(str(excel_file)))

        assert len(result) == 2

        assert result[0] == (2, {"id": 1, "name": "admin", "role": "superuser"})
        assert result[1] == (3, {"id": 2, "name": "user", "role": "manager"})

    def test_parse_excel_generator_header_sanitization(self, tmp_path):
        excel_file = tmp_path / "test_headers.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active

        ws.append([" First Name ", "Last\nName", None, "Age"])
        ws.append(["John", "Doe", "secret_data", 30])
        wb.save(excel_file)

        result = list(parse_excel_generator(str(excel_file)))

        assert len(result) == 1

        row_idx, row_data = result[0]
        assert row_idx == 2
        assert row_data["FirstName"] == "John"
        assert row_data["LastName"] == "Doe"
        assert row_data["col_2"] == "secret_data"
        assert row_data["Age"] == 30

    def test_parse_excel_generator_skips_empty_rows(self, tmp_path):
        excel_file = tmp_path / "test_empty.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active

        ws.append(["item", "price"])
        ws.append(["apple", 100])
        ws.append([None, None])
        ws.append(["banana", 50])
        wb.save(excel_file)

        result = list(parse_excel_generator(str(excel_file)))

        assert len(result) == 2
        assert result[0] == (2, {"item": "apple", "price": 100})
        assert result[1] == (4, {"item": "banana", "price": 50})
