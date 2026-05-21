import tempfile
import uuid
from pathlib import Path
import openpyxl

from app.api.v1.schemas.task import BatchGenerateData, ProductsReportData, StatsReportData


def generate_batch_excel_report(
        batch: BatchGenerateData,
        products: ProductsReportData,
        stats: StatsReportData,
) -> str:
    wb = openpyxl.Workbook()

    # list 1
    ws_info = wb.active
    ws_info.title = "Информация"

    ws_info.append(["Номер партии:", batch.batch_number])
    ws_info.append(["Дата партии:", batch.batch_date.strftime("%d-%m-%Y")])
    ws_info.append(["Статус:", "Закрыта" if batch.is_closed else "Открыта"])
    ws_info.append(["Рабочий центр:", f"Цех №{batch.work_center_id}"])
    ws_info.append(["Смена:", f"{batch.shift} смена"])
    ws_info.append(["Бригада:", f"Бригада {batch.team}"])
    ws_info.append(["Номенклатура:", batch.nomenclature])
    ws_info.append(["Время смены:", f"{batch.shift_start.strftime('%H:%M')} - {batch.shift_end.strftime('%H:%M')}"])

    # list 2
    ws_products = wb.create_sheet(title="Продукция")

    ws_products.append(["ID", "Уникальный код", "Статус", "Время агрегации"])
    for p in products.products:
        ws_products.append([
            p.id,
            p.unique_code,
            "Да" if p.is_aggregated else "Нет",
            p.aggregated_at.strftime("%d-%m-%Y %H:%M:%S") if p.aggregated_at else "-"
        ])

    # list 3
    ws_stats = wb.create_sheet(title="Статистика")

    ws_stats.append(["Всего продукции:", stats.total_products])
    ws_stats.append(["Агрегировано:", stats.aggregated_products])
    ws_stats.append(["Осталось:", stats.unaggregated_products])
    ws_stats.append(["Процент выполнения:", f"{stats.percent_aggregated}%"])
    ws_stats.append(["Средняя скорость:", f"{stats.avg_speed} ед/час" if stats.avg_speed else "-"])

    file_path = Path(tempfile.gettempdir()) / f"batch_{batch.batch_number}_report_{uuid.uuid4().hex}.xlsx"
    wb.save(file_path)

    return str(file_path)


def generate_batches_excel(batches: list[BatchGenerateData]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Экспорт партий"

    headers = [
        "Номер партии", "Дата партии", "Статус", "Рабочий центр",
        "Смена", "Бригада", "Номенклатура", "Время смены"
    ]
    ws.append(headers)

    for batch in batches:
        ws.append([
            batch.batch_number,
            batch.batch_date.strftime("%d.%m.%Y"),
            "Закрыта" if batch.is_closed else "Открыта",
            f"Цех №{batch.work_center_id}",
            f"{batch.shift} смена",
            f"Бригада {batch.team}",
            batch.nomenclature,
            f"{batch.shift_start.strftime('%H:%M')} - {batch.shift_end.strftime('%H:%M')}"
        ])


    unique_filename = f"export_batches_{uuid.uuid4().hex}.xlsx"
    file_path = Path(tempfile.gettempdir()) / unique_filename

    wb.save(file_path)

    return str(file_path)